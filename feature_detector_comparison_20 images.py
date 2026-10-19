"""2D Feature Detector Comparison: MSER, SIFT, FAST, Harris, ORB  (v3, revised)

Usage:
    python feature_detector_comparison_v3.py            # folder picker dialog
    python feature_detector_comparison_v3.py <folder>   # headless

Optional inputs in the image folder
-----------------------------------
* masks/<image_stem>.png   -> object mask (non-zero = object). When present,
  keypoints, coverage, density and NN distance are restricted to the object.
* ground_truth_homographies.csv -> Image_A, Image_B, H00..H22 (A pixels -> B
  pixels, in ORIGINAL image coordinates; rescaling is handled internally).
  Only valid for planar scenes / pure-rotation pairs.

What changed versus v2
----------------------
1. Matching descriptors use ONE policy for every detector (fixed size, upright,
   octave 0) so the match rate no longer rewards SIFT/ORB for supplying their
   own scale/orientation. MATCH_DESCRIPTOR_POLICY = "native" restores the old
   behaviour as a sensitivity check.
2. Failed pairs are no longer dropped: success rate, failures-as-zero means,
   success-only means and standard deviations are all reported.
3. Stricter geometry: >= 15 mutual matches, success needs >= 12 RANSAC inliers,
   seeded RANSAC, inlier ratio and absolute inlier counts reported.
4. Timing is split into core detector time and Python keypoint-conversion time
   (relevant for MSER and Harris), with a warm-up call.
5. Order is now: rank by response -> spatial suppression -> cap at N. Every
   detector therefore gets the same final keypoint budget of distinct points.
6. Settings equalised as far as is meaningful: optional common resize, MSER
   area limits scaled to image size, thresholds low enough that the cap binds.
   "Raw keypoints" is informational only and is no longer charted.
7. Optional object masks for coverage / density / NN distance; NN distance is
   normalised by sqrt(region area).
Minor: unused psutil removed, unicode-safe image I/O, std devs everywhere,
descriptors timed on the same final keypoints used for matching.

Remaining limitations (inherent, not bugs)
------------------------------------------
* Fundamental-matrix inliers are still a proxy: a rigid object on a static
  background violates a single F, and epipolar-line outliers can pass.
  True repeatability needs ground-truth geometry (supplied homographies).
* MSER regions are represented by centroids ranked by area (custom adapter).
* Timing is a single-machine benchmark.

Dependencies:  pip install opencv-python numpy pandas matplotlib openpyxl
"""

import itertools
import math
import os
import sys
import time
import traceback

import cv2
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ALGORITHMS = ["MSER", "SIFT", "FAST", "HARRIS", "ORB"]
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp")

# ---- selection / budget -------------------------------------------------
FINAL_KEYPOINT_BUDGET = 1000      # applied AFTER spatial suppression
SUPPRESSION_RADIUS_PX = 3.0
RESIZE_MAX_SIDE = 1280            # None = keep original resolution
COVERAGE_RADIUS_PX = 5
MASK_FOLDER_NAME = "masks"

# ---- detector settings (low thresholds so the budget cap is what binds) --
HARRIS_RELATIVE_THRESHOLD = 0.001
FAST_THRESHOLD = 10
ORB_MAX_FEATURES = 10000
SIFT_MAX_FEATURES = 0
MSER_REFERENCE_PIXELS = 640 * 480  # original 30 / 14400 limits were tuned for this
MSER_MIN_AREA_REF = 30
MSER_MAX_AREA_REF = 14400

# ---- timing --------------------------------------------------------------
TIMING_RUNS = 5                   # after one warm-up call

# ---- matching proxy ------------------------------------------------------
MATCH_DESCRIPTOR_POLICY = "fixed"  # "fixed" (fair) or "native" (sensitivity)
MATCH_KP_SIZE = 16.0
MATCH_RATIO = 0.75
MIN_GEOMETRIC_MATCHES = 15        # mutual matches needed to attempt RANSAC
MIN_INLIERS_SUCCESS = 12          # inliers needed to call a pair successful
FUNDAMENTAL_RANSAC_THRESHOLD_PX = 1.5
RANSAC_SEED = 12345

# ---- homography repeatability -------------------------------------------
REPEATABILITY_DISTANCE_PX = 3.0
HOMOGRAPHY_FILE_NAME = "ground_truth_homographies.csv"

COLORS = {  # BGR
    "MSER": (255, 80, 80), "SIFT": (60, 210, 80), "FAST": (50, 130, 255),
    "HARRIS": (240, 190, 30), "ORB": (210, 70, 230),
}


# =========================================================================
# I/O helpers
# =========================================================================
def choose_input_folder():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    try:
        return filedialog.askdirectory(
            title="Select folder containing images of the same model")
    finally:
        root.destroy()


def list_images(folder):
    return sorted(
        os.path.join(folder, n) for n in os.listdir(folder)
        if n.lower().endswith(IMAGE_EXTENSIONS)
        and os.path.isfile(os.path.join(folder, n)))


def read_image(path, flags=cv2.IMREAD_COLOR):
    """Unicode-safe imread (cv2.imread fails on non-ASCII Windows paths)."""
    try:
        data = np.fromfile(path, dtype=np.uint8)
        return cv2.imdecode(data, flags) if data.size else None
    except OSError:
        return None


def write_image(path, image):
    ext = os.path.splitext(path)[1] or ".jpg"
    ok, buf = cv2.imencode(ext, image)
    if ok:
        buf.tofile(path)
    return bool(ok)


def load_mask(folder, filename, shape):
    path = os.path.join(folder, MASK_FOLDER_NAME, os.path.splitext(filename)[0] + ".png")
    if not os.path.isfile(path):
        return None
    mask = read_image(path, cv2.IMREAD_GRAYSCALE)
    if mask is None:
        return None
    mask = cv2.resize(mask, (shape[1], shape[0]), interpolation=cv2.INTER_NEAREST)
    mask = mask > 0
    return mask if mask.any() else None


# =========================================================================
# Detectors (each returns keypoints, core_ms, conversion_ms)
# =========================================================================
_mser_cache = {}


def get_mser(height, width):
    key = (height, width)
    if key not in _mser_cache:
        scale = (height * width) / MSER_REFERENCE_PIXELS
        _mser_cache[key] = cv2.MSER_create(
            5,
            max(10, int(round(MSER_MIN_AREA_REF * scale))),
            max(100, int(round(MSER_MAX_AREA_REF * scale))),
            0.25, 0.20)
    return _mser_cache[key]


def make_detectors():
    return {
        "SIFT": cv2.SIFT_create(nfeatures=SIFT_MAX_FEATURES),
        "FAST": cv2.FastFeatureDetector_create(
            threshold=FAST_THRESHOLD, nonmaxSuppression=True),
        "ORB": cv2.ORB_create(nfeatures=ORB_MAX_FEATURES),
    }


def detect_mser(gray, detectors):
    detector = get_mser(*gray.shape)
    t0 = time.perf_counter()
    regions, _ = detector.detectRegions(gray)
    t1 = time.perf_counter()
    keypoints = []
    for region in regions:
        if region is None or len(region) == 0:
            continue
        x, y = np.mean(region, axis=0)
        area = float(len(region))
        keypoints.append(cv2.KeyPoint(
            float(x), float(y), float(max(3.0, math.sqrt(area))), -1, area))
    t2 = time.perf_counter()
    return keypoints, (t1 - t0) * 1e3, (t2 - t1) * 1e3


def detect_harris(gray, detectors):
    t0 = time.perf_counter()
    response = cv2.cornerHarris(np.float32(gray), blockSize=2, ksize=3, k=0.04)
    max_response = float(response.max()) if response.size else 0.0
    if not math.isfinite(max_response) or max_response <= 0:
        return [], (time.perf_counter() - t0) * 1e3, 0.0
    local_max = response == cv2.dilate(response, np.ones((3, 3), np.uint8))
    mask = local_max & (response > HARRIS_RELATIVE_THRESHOLD * max_response)
    t1 = time.perf_counter()
    ys, xs = np.nonzero(mask)
    keypoints = [cv2.KeyPoint(float(x), float(y), 5.0, -1, float(response[y, x]))
                 for y, x in zip(ys, xs)]
    t2 = time.perf_counter()
    return keypoints, (t1 - t0) * 1e3, (t2 - t1) * 1e3


def detect_native(name):
    def run(gray, detectors):
        t0 = time.perf_counter()
        kps = detectors[name].detect(gray, None) or []
        return list(kps), (time.perf_counter() - t0) * 1e3, 0.0
    return run


DETECT_FUNCS = {
    "MSER": detect_mser, "HARRIS": detect_harris,
    "SIFT": detect_native("SIFT"), "FAST": detect_native("FAST"),
    "ORB": detect_native("ORB"),
}


def detect_all(gray, detectors):
    results, timing = {}, {}
    for name in ALGORITHMS:
        func = DETECT_FUNCS[name]
        func(gray, detectors)  # warm-up, discarded
        core, conv, keypoints = [], [], []
        for _ in range(TIMING_RUNS):
            keypoints, c, v = func(gray, detectors)
            core.append(c)
            conv.append(v)
        results[name] = sorted(keypoints, key=lambda k: float(k.response), reverse=True)
        timing[name] = (float(np.mean(core)), float(np.mean(conv)))
    return results, timing


# =========================================================================
# Selection: mask -> suppression -> cap
# =========================================================================
def filter_by_mask(keypoints, mask):
    if mask is None:
        return keypoints
    h, w = mask.shape
    out = []
    for kp in keypoints:
        x = int(np.clip(round(kp.pt[0]), 0, w - 1))
        y = int(np.clip(round(kp.pt[1]), 0, h - 1))
        if mask[y, x]:
            out.append(kp)
    return out


def spatial_suppression(keypoints, radius=SUPPRESSION_RADIUS_PX):
    """Greedy response-ranked suppression using a grid hash (O(n))."""
    cell = max(radius, 1e-6)
    radius_sq = radius * radius
    grid, kept, removed = {}, [], 0
    for kp in keypoints:  # already sorted by response, descending
        x, y = kp.pt
        cx, cy = int(x // cell), int(y // cell)
        clash = False
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                for ox, oy in grid.get((cx + dx, cy + dy), ()):
                    if (ox - x) ** 2 + (oy - y) ** 2 <= radius_sq:
                        clash = True
                        break
                if clash:
                    break
            if clash:
                break
        if clash:
            removed += 1
        else:
            kept.append(kp)
            grid.setdefault((cx, cy), []).append((x, y))
    return kept, removed


# =========================================================================
# Spatial metrics
# =========================================================================
def coverage_percent(keypoints, shape, region_mask, radius=COVERAGE_RADIUS_PX):
    height, width = shape
    region = region_mask if region_mask is not None else np.ones(shape, bool)
    region_area = int(region.sum())
    if not keypoints or region_area == 0:
        return 0.0
    canvas = np.zeros(shape, np.uint8)
    for kp in keypoints:
        x = int(np.clip(round(kp.pt[0]), 0, width - 1))
        y = int(np.clip(round(kp.pt[1]), 0, height - 1))
        cv2.circle(canvas, (x, y), int(radius), 1, thickness=-1)
    return 100.0 * np.count_nonzero((canvas > 0) & region) / region_area


def mean_nn_distance(keypoints):
    if len(keypoints) < 2:
        return float("nan")
    pts = np.asarray([kp.pt for kp in keypoints], np.float32)
    nearest = []
    for start in range(0, len(pts), 500):
        chunk = pts[start:start + 500]
        d2 = np.sum((chunk[:, None, :] - pts[None, :, :]) ** 2, axis=2)
        d2[np.arange(len(chunk)), np.arange(start, start + len(chunk))] = np.inf
        nearest.extend(np.sqrt(d2.min(axis=1)).tolist())
    return float(np.mean(nearest))


# =========================================================================
# Cross-view matching proxy
# =========================================================================
def clone_keypoints(keypoints):
    return [cv2.KeyPoint(float(k.pt[0]), float(k.pt[1]), float(k.size),
                         float(k.angle), float(k.response), int(k.octave),
                         int(k.class_id)) for k in keypoints]


def normalise_for_matching(keypoints):
    """Give every detector the same descriptor geometry (fair comparison)."""
    if MATCH_DESCRIPTOR_POLICY == "native":
        return clone_keypoints(keypoints)
    return [cv2.KeyPoint(float(k.pt[0]), float(k.pt[1]), MATCH_KP_SIZE,
                         0.0, float(k.response), 0, -1) for k in keypoints]


def compute_matching_descriptors(gray, keypoints, extractor):
    if not keypoints:
        return [], None
    kps, desc = extractor.compute(gray, normalise_for_matching(keypoints))
    if desc is None or kps is None or len(kps) == 0:
        return [], None
    return list(kps), desc


def mutual_ratio_matches(desc_a, desc_b, ratio=MATCH_RATIO):
    if desc_a is None or desc_b is None or len(desc_a) < 2 or len(desc_b) < 2:
        return []
    matcher = cv2.BFMatcher(cv2.NORM_L2)

    def one_way(a, b):
        good = {}
        for pair in matcher.knnMatch(a, b, k=2):
            if len(pair) == 2 and pair[0].distance < ratio * pair[1].distance:
                good[pair[0].queryIdx] = pair[0].trainIdx
        return good

    ab, ba = one_way(desc_a, desc_b), one_way(desc_b, desc_a)
    return [(i, j) for i, j in ab.items() if ba.get(j) == i]


def geometric_match(view_a, view_b):
    """Returns dict: mutual, inliers, rate, inlier_ratio, success."""
    out = {"mutual": 0, "inliers": 0, "rate": float("nan"),
           "inlier_ratio": float("nan"), "success": False}
    kp_a, desc_a = view_a["match_kp"], view_a["match_desc"]
    kp_b, desc_b = view_b["match_kp"], view_b["match_desc"]
    if desc_a is None or desc_b is None:
        return out
    matches = mutual_ratio_matches(desc_a, desc_b)
    out["mutual"] = len(matches)
    if len(matches) < MIN_GEOMETRIC_MATCHES:
        return out
    pts_a = np.float32([kp_a[i].pt for i, _ in matches]).reshape(-1, 1, 2)
    pts_b = np.float32([kp_b[j].pt for _, j in matches]).reshape(-1, 1, 2)
    try:
        cv2.setRNGSeed(RANSAC_SEED)  # reproducible RANSAC
        _, mask = cv2.findFundamentalMat(
            pts_a, pts_b, cv2.FM_RANSAC, FUNDAMENTAL_RANSAC_THRESHOLD_PX, 0.99)
    except cv2.error:
        return out
    if mask is None or mask.size != len(matches):
        return out
    inliers = int(np.count_nonzero(mask.ravel()))
    denominator = max(1, min(len(kp_a), len(kp_b)))
    out.update(inliers=inliers,
               rate=min(100.0, 100.0 * inliers / denominator),
               inlier_ratio=100.0 * inliers / len(matches),
               success=inliers >= MIN_INLIERS_SUCCESS)
    return out


# =========================================================================
# Optional ground-truth homography repeatability
# =========================================================================
def load_ground_truth_homographies(folder):
    path = os.path.join(folder, HOMOGRAPHY_FILE_NAME)
    if not os.path.isfile(path):
        return {}, path
    table = pd.read_csv(path)
    required = ["Image_A", "Image_B"] + [f"H{r}{c}" for r in range(3) for c in range(3)]
    missing = [c for c in required if c not in table.columns]
    if missing:
        raise ValueError(f"{HOMOGRAPHY_FILE_NAME} missing columns: {', '.join(missing)}")
    result = {}
    for _, row in table.iterrows():
        a, b = str(row["Image_A"]), str(row["Image_B"])
        m = np.asarray([[row[f"H{r}{c}"] for c in range(3)] for r in range(3)], np.float64)
        if not np.all(np.isfinite(m)) or abs(np.linalg.det(m)) < 1e-12:
            print(f"WARNING: ignoring invalid homography {a} -> {b}")
            continue
        result[(a, b)] = m
    return result, path


def get_pair_homography(homographies, a, b):
    m = homographies.get((a, b))
    if m is not None:
        return m
    rev = homographies.get((b, a))
    if rev is None:
        return None
    try:
        return np.linalg.inv(rev)
    except np.linalg.LinAlgError:
        return None


def homography_repeatability(view_a, view_b, homography_original):
    """Homography is in ORIGINAL pixel coords; convert to the resized frames."""
    sa, sb = view_a["scale"], view_b["scale"]
    homography = np.diag([sb, sb, 1.0]) @ homography_original @ np.diag([1 / sa, 1 / sa, 1.0])
    ka, kb = view_a["detected_kp"], view_b["detected_kp"]
    if not ka or not kb:
        return float("nan"), 0, 0
    try:
        inverse = np.linalg.inv(homography)
    except np.linalg.LinAlgError:
        return float("nan"), 0, 0
    pa = np.asarray([k.pt for k in ka], np.float32)
    pb = np.asarray([k.pt for k in kb], np.float32)
    proj_a = cv2.perspectiveTransform(pa.reshape(-1, 1, 2), homography).reshape(-1, 2)
    proj_b = cv2.perspectiveTransform(pb.reshape(-1, 1, 2), inverse).reshape(-1, 2)

    def inside(p, w, h):
        return (np.isfinite(p).all(axis=1) & (p[:, 0] >= 0) & (p[:, 0] < w)
                & (p[:, 1] >= 0) & (p[:, 1] < h))

    ids_a = np.flatnonzero(inside(proj_a, view_b["width"], view_b["height"]))
    ids_b = np.flatnonzero(inside(proj_b, view_a["width"], view_a["height"]))
    denominator = min(len(ids_a), len(ids_b))
    if denominator == 0:
        return float("nan"), 0, 0
    src = np.ascontiguousarray(proj_a[ids_a], np.float32)
    dst = np.ascontiguousarray(pb[ids_b], np.float32)
    matcher = cv2.BFMatcher(cv2.NORM_L2)
    ab = matcher.match(src, dst)
    ba = {m.queryIdx: m.trainIdx for m in matcher.match(dst, src)}
    repeated = sum(1 for m in ab
                   if m.distance <= REPEATABILITY_DISTANCE_PX and ba.get(m.trainIdx) == m.queryIdx)
    return min(100.0, 100.0 * repeated / denominator), repeated, denominator


# =========================================================================
# Native descriptor timing (on the same final keypoints used elsewhere)
# =========================================================================
def native_descriptor_stats(name, gray, selected, detectors):
    if name not in ("SIFT", "ORB") or not selected:
        return float("nan"), 0
    durations, count = [], 0
    detectors[name].compute(gray, clone_keypoints(selected))  # warm-up
    for _ in range(TIMING_RUNS):
        kps = clone_keypoints(selected)
        t0 = time.perf_counter()
        _, desc = detectors[name].compute(gray, kps)
        durations.append((time.perf_counter() - t0) * 1e3)
        count = 0 if desc is None else len(desc)
    return float(np.mean(durations)), count


# =========================================================================
# Reporting
# =========================================================================
def save_chart(summary, column, ylabel, title, filename, folder, std_column=None):
    labels, values, errors = [], [], []
    for _, row in summary.iterrows():
        v = pd.to_numeric(pd.Series([row.get(column)]), errors="coerce").iloc[0]
        if pd.notna(v) and np.isfinite(v):
            labels.append(str(row["Algorithm"]))
            values.append(float(v))
            e = (pd.to_numeric(pd.Series([row.get(std_column)]), errors="coerce").iloc[0]
                 if std_column else np.nan)
            errors.append(float(e) if pd.notna(e) else 0.0)
    fig, ax = plt.subplots(figsize=(8, 5))
    if labels:
        ax.bar(labels, values, yerr=errors if std_column else None, capsize=4)
    else:
        ax.text(0.5, 0.5, "No valid data", ha="center", va="center", transform=ax.transAxes)
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    fig.tight_layout()
    fig.savefig(os.path.join(folder, filename), dpi=250)
    plt.close(fig)


def style_workbook(writer):
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and ws.max_column > 0:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ws.columns:
            longest = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 14), 42)
        ws.row_dimensions[1].height = 32


def stat(series):
    s = pd.to_numeric(series, errors="coerce").dropna()
    return (float(s.mean()), float(s.std(ddof=1)) if len(s) > 1 else float("nan")) if len(s) else (float("nan"), float("nan"))


def r(value, digits=2):
    return round(value, digits) if isinstance(value, (int, float)) and math.isfinite(value) else float("nan")


# =========================================================================
# Main
# =========================================================================
def main(folder=None, interactive=True):
    if folder is None:
        folder = choose_input_folder()
    if not folder:
        print("Cancelled: no folder selected.")
        return
    image_paths = list_images(folder)
    if not image_paths:
        raise RuntimeError("No supported image files found in the selected folder.")
    if interactive and len(image_paths) != 20:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        try:
            if not messagebox.askyesno(
                    "Image count",
                    f"Found {len(image_paths)} image(s); intended setup is 20.\nContinue?"):
                return
        finally:
            root.destroy()

    output_folder = os.path.join(folder, "feature_detection_analysis_2d_v3")
    marked_folder = os.path.join(output_folder, "marked_images")
    os.makedirs(marked_folder, exist_ok=True)

    homographies, h_path = load_ground_truth_homographies(folder)
    print(f"Loaded {len(homographies)} homography pair(s)." if homographies else
          "No ground-truth homographies: repeatability will be N/A (matching stays a proxy).")

    detectors = make_detectors()
    sift_matcher_extractor = cv2.SIFT_create()
    per_rows = []
    views = {n: [] for n in ALGORITHMS}
    masks_used = 0
    t_total = time.perf_counter()

    for idx, path in enumerate(image_paths, 1):
        filename = os.path.basename(path)
        print(f"[{idx}/{len(image_paths)}] {filename}")
        image = read_image(path)
        if image is None:
            print("  WARNING: unreadable, skipping.")
            continue
        scale = 1.0
        if RESIZE_MAX_SIDE and max(image.shape[:2]) > RESIZE_MAX_SIDE:
            scale = RESIZE_MAX_SIDE / max(image.shape[:2])
            image = cv2.resize(image, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        height, width = gray.shape
        mask = load_mask(folder, filename, gray.shape)
        masks_used += mask is not None
        region_area = int(mask.sum()) if mask is not None else width * height

        detected, timing = detect_all(gray, detectors)

        for name in ALGORITHMS:
            raw = detected[name]
            in_region = filter_by_mask(raw, mask)
            suppressed, removed = spatial_suppression(in_region)
            selected = suppressed[:FINAL_KEYPOINT_BUDGET]

            desc_ms, desc_count = native_descriptor_stats(name, gray, selected, detectors)
            match_kp, match_desc = compute_matching_descriptors(
                gray, selected, sift_matcher_extractor)

            nn_px = mean_nn_distance(selected)
            nn_pct = 100.0 * nn_px / math.sqrt(region_area) if math.isfinite(nn_px) else float("nan")
            core_ms, conv_ms = timing[name]

            per_rows.append({
                "Image": filename, "Algorithm": name,
                "Width_px": width, "Height_px": height, "Resize_Scale": round(scale, 5),
                "Object_Mask_Used": mask is not None,
                "Region_Area_px": region_area,
                "Detection_Core_Time_ms": core_ms,
                "Keypoint_Conversion_Time_ms": conv_ms,
                "Detection_Total_Time_ms": core_ms + conv_ms,
                "Native_Descriptor_Time_ms": desc_ms,
                "Raw_Keypoints_Informational": len(raw),
                "Raw_In_Region": len(in_region),
                "Points_Removed_By_Suppression": removed,
                "Suppression_Rate_%": 100.0 * removed / len(in_region) if in_region else 0.0,
                "Distinct_Keypoints_Available": len(suppressed),
                "Final_Selected_Keypoints": len(selected),
                "Budget_Filled": len(selected) >= FINAL_KEYPOINT_BUDGET,
                "Mean_NN_Distance_px": nn_px,
                "Mean_NN_Distance_%_sqrt_Region_Area": nn_pct,
                "Feature_Density_per_Megapixel": len(selected) / max(1, region_area) * 1e6,
                "2D_Coverage_%_of_Region": coverage_percent(selected, gray.shape, mask),
                "Native_Descriptor_Count": desc_count if name in ("SIFT", "ORB") else float("nan"),
            })
            views[name].append({
                "filename": filename, "scale": scale, "width": width, "height": height,
                "detected_kp": selected, "match_kp": match_kp, "match_desc": match_desc,
            })
            marked = cv2.drawKeypoints(
                image, selected, None, color=COLORS[name],
                flags=cv2.DRAW_MATCHES_FLAGS_DRAW_RICH_KEYPOINTS)
            write_image(os.path.join(
                marked_folder, f"{os.path.splitext(filename)[0]}_{name.lower()}_keypoints.jpg"), marked)

    n_processed = len(views[ALGORITHMS[0]])
    if n_processed < 1:
        raise RuntimeError("None of the image files could be read.")

    # ---- all unordered pairs -------------------------------------------
    pair_rows = []
    for name in ALGORITHMS:
        for a, b in itertools.combinations(views[name], 2):
            g = geometric_match(a, b)
            H = get_pair_homography(homographies, a["filename"], b["filename"])
            rep = homography_repeatability(a, b, H) if H is not None else (float("nan"), 0, 0)
            pair_rows.append({
                "Algorithm": name, "Image_A": a["filename"], "Image_B": b["filename"],
                "Mutual_Ratio_Test_Matches": g["mutual"],
                "RANSAC_Inliers": g["inliers"],
                "Pair_Success": g["success"],
                "Inlier_Ratio_%_of_Mutual": g["inlier_ratio"],
                "Match_Rate_%_Estimated_Only": g["rate"],
                "Match_Rate_%_Failures_As_Zero": g["rate"] if g["success"] else 0.0,
                "Ground_Truth_Homography_Available": H is not None,
                "Homography_Repeatability_%": rep[0],
                "Homography_Repeated_Keypoints": rep[1] if H is not None else float("nan"),
                "Homography_Repeatability_Denominator": rep[2] if H is not None else float("nan"),
            })

    per_df = pd.DataFrame(per_rows)
    pair_df = pd.DataFrame(pair_rows)

    # ---- summary -------------------------------------------------------
    summary_rows = []
    for name in ALGORITHMS:
        s = per_df[per_df["Algorithm"] == name]
        p = pair_df[pair_df["Algorithm"] == name] if not pair_df.empty else pd.DataFrame()
        row = {"Algorithm": name, "Images_Processed": len(s)}
        for col, label in [("Detection_Core_Time_ms", "Detection_Core_Time_ms"),
                           ("Keypoint_Conversion_Time_ms", "Conversion_Time_ms"),
                           ("Detection_Total_Time_ms", "Detection_Total_Time_ms")]:
            m, sd = stat(s[col])
            row[f"Mean_{label}_per_Image"], row[f"Std_{label}"] = r(m, 4), r(sd, 4)
        m, sd = stat(s["Native_Descriptor_Time_ms"])
        row["Mean_Native_Descriptor_Time_ms"], row["Std_Native_Descriptor_Time_ms"] = r(m, 4), r(sd, 4)
        row["Total_Raw_Keypoints_Informational"] = int(s["Raw_Keypoints_Informational"].sum())
        row["Overall_Suppression_Rate_%"] = r(
            100.0 * s["Points_Removed_By_Suppression"].sum() / max(1, s["Raw_In_Region"].sum()))
        row["Mean_Distinct_Keypoints_Available"] = r(stat(s["Distinct_Keypoints_Available"])[0], 1)
        row["Mean_Final_Selected_Keypoints"] = r(stat(s["Final_Selected_Keypoints"])[0], 1)
        row["Images_Where_Budget_Filled_%"] = r(100.0 * s["Budget_Filled"].mean())
        for col, label in [("Mean_NN_Distance_%_sqrt_Region_Area", "NN_Distance_%_sqrt_Region_Area"),
                           ("Feature_Density_per_Megapixel", "Feature_Density_per_Megapixel"),
                           ("2D_Coverage_%_of_Region", "2D_Coverage_%_of_Region")]:
            m, sd = stat(s[col])
            row[f"Mean_{label}"], row[f"Std_{label}"] = r(m, 4), r(sd, 4)

        n_pairs = len(p)
        row["Total_Pairs"] = n_pairs
        if n_pairs:
            row["Pair_Success_Rate_%"] = r(100.0 * p["Pair_Success"].mean())
            m, sd = stat(p["Match_Rate_%_Failures_As_Zero"])
            row["Mean_Match_Rate_%_Failures_As_Zero"], row["Std_Match_Rate_%_Failures_As_Zero"] = r(m), r(sd)
            ok = p[p["Pair_Success"]]
            m, sd = stat(ok["Match_Rate_%_Estimated_Only"])
            row["Mean_Match_Rate_%_Successful_Pairs"], row["Std_Match_Rate_%_Successful_Pairs"] = r(m), r(sd)
            row["Mean_Inliers_All_Pairs_Failures_As_Zero"] = r(
                float(np.where(p["Pair_Success"], p["RANSAC_Inliers"], 0).mean()), 1)
            row["Mean_Inlier_Ratio_%_Successful_Pairs"] = r(stat(ok["Inlier_Ratio_%_of_Mutual"])[0])
            hv = p["Homography_Repeatability_%"].dropna()
            m, sd = stat(hv)
            row["Valid_Homography_Pairs"] = int(len(hv))
            row["Mean_Homography_Repeatability_%"], row["Std_Homography_Repeatability_%"] = r(m), r(sd)
        summary_df = None
        summary_rows.append(row)
    summary_df = pd.DataFrame(summary_rows)

    # ---- write ----------------------------------------------------------
    paths = {
        "per": os.path.join(output_folder, "per_image_metrics.csv"),
        "sum": os.path.join(output_folder, "comparison_table.csv"),
        "pair": os.path.join(output_folder, "cross_view_matching.csv"),
        "xlsx": os.path.join(output_folder, "comparison_table.xlsx"),
    }

    def out(df):
        return df.round(4).astype(object).where(df.notna(), "N/A")

    out(per_df).to_csv(paths["per"], index=False)
    out(summary_df).to_csv(paths["sum"], index=False)
    out(pair_df).to_csv(paths["pair"], index=False)

    settings = pd.DataFrame([
        ["Input folder", folder],
        ["Images processed", n_processed],
        ["Object masks used (images)", masks_used],
        ["Resize", f"Long side <= {RESIZE_MAX_SIDE}px (INTER_AREA)" if RESIZE_MAX_SIDE else "None"],
        ["Pipeline", "Detect -> rank by response -> object-mask filter -> spatial suppression -> cap"],
        ["Final keypoint budget (after suppression)", FINAL_KEYPOINT_BUDGET],
        ["Suppression radius (px)", SUPPRESSION_RADIUS_PX],
        ["MSER", f"Centroids ranked by area; area limits {MSER_MIN_AREA_REF}/{MSER_MAX_AREA_REF} scaled from {MSER_REFERENCE_PIXELS} px"],
        ["Harris relative threshold", HARRIS_RELATIVE_THRESHOLD],
        ["FAST threshold", FAST_THRESHOLD],
        ["ORB nfeatures", ORB_MAX_FEATURES],
        ["SIFT nfeatures", SIFT_MAX_FEATURES],
        ["Timing", f"1 warm-up + mean of {TIMING_RUNS} calls; core vs Python conversion reported separately"],
        ["Matching descriptor policy", f"{MATCH_DESCRIPTOR_POLICY} (fixed = size {MATCH_KP_SIZE}, upright, octave 0 for ALL detectors)"],
        ["Matching", f"Mutual ratio test ({MATCH_RATIO}); needs >= {MIN_GEOMETRIC_MATCHES} mutual matches; seeded RANSAC fundamental matrix ({FUNDAMENTAL_RANSAC_THRESHOLD_PX}px)"],
        ["Pair success", f">= {MIN_INLIERS_SUCCESS} RANSAC inliers; failed pairs counted as 0 in 'Failures_As_Zero' metrics"],
        ["Match rate", "Inliers / min(descriptor count of the two views). Proxy only, NOT repeatability"],
        ["Homography repeatability", f"Needs {HOMOGRAPHY_FILE_NAME}; planar scenes only; mutual NN within {REPEATABILITY_DISTANCE_PX}px; homographies rescaled for resizing"],
        ["Coverage / density / NN", "Computed inside object mask if available, else whole image; NN normalised by sqrt(region area); 2D only"],
        ["Raw keypoints", "Informational only; reflects detector parameters, not quality"],
        ["Memory", "Not ranked (shared-process RSS cannot isolate algorithms)"],
        ["Std columns", "Sample std over images (or pairs); pairs share images so are not independent"],
    ], columns=["Setting", "Value"])

    with pd.ExcelWriter(paths["xlsx"], engine="openpyxl") as writer:
        out(summary_df).to_excel(writer, sheet_name="Comparison", index=False)
        out(per_df).to_excel(writer, sheet_name="Per Image", index=False)
        out(pair_df).to_excel(writer, sheet_name="Cross-view Pairs", index=False)
        settings.to_excel(writer, sheet_name="Methodology", index=False)
        style_workbook(writer)

    charts = [
        ("Mean_Detection_Core_Time_ms_per_Image", "ms", "Core Detection Time per Image", "core_detection_time.png", "Std_Detection_Core_Time_ms"),
        ("Mean_Detection_Total_Time_ms_per_Image", "ms", "Total Detection Time (incl. Python conversion)", "total_detection_time.png", "Std_Detection_Total_Time_ms"),
        ("Mean_Final_Selected_Keypoints", "Count", "Final Selected Keypoints (budget %d)" % FINAL_KEYPOINT_BUDGET, "final_keypoints.png", None),
        ("Overall_Suppression_Rate_%", "Percent", "Spatial Suppression Rate", "suppression_rate.png", None),
        ("Pair_Success_Rate_%", "Percent", "Cross-view Pair Success Rate", "pair_success_rate.png", None),
        ("Mean_Match_Rate_%_Failures_As_Zero", "Percent", "Match Rate, Failures = 0 (Proxy)", "match_rate_failures_zero.png", "Std_Match_Rate_%_Failures_As_Zero"),
        ("Mean_Homography_Repeatability_%", "Percent", "Homography Repeatability (valid GT pairs)", "homography_repeatability.png", "Std_Homography_Repeatability_%"),
        ("Mean_NN_Distance_%_sqrt_Region_Area", "% of sqrt(area)", "Normalised Mean NN Distance", "nn_distance.png", "Std_NN_Distance_%_sqrt_Region_Area"),
        ("Mean_2D_Coverage_%_of_Region", "Percent", "2D Coverage of Region", "coverage.png", "Std_2D_Coverage_%_of_Region"),
    ]
    for col, ylabel, title, fname, std_col in charts:
        if col in summary_df.columns:
            save_chart(summary_df, col, ylabel, title, fname, output_folder, std_col)

    print("\nCOMPARISON TABLE")
    print(summary_df.T.to_string(header=False))
    print(f"\nDone in {time.perf_counter() - t_total:.1f}s -> {output_folder}")

    if interactive:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        try:
            messagebox.showinfo("Analysis complete", f"Results saved in:\n{output_folder}")
        finally:
            root.destroy()


if __name__ == "__main__":
    try:
        if len(sys.argv) > 1:
            main(sys.argv[1], interactive=False)
        else:
            main()
    except Exception as exc:
        traceback.print_exc()
        if len(sys.argv) <= 1:
            try:
                import tkinter as tk
                from tkinter import messagebox
                root = tk.Tk()
                root.withdraw()
                messagebox.showerror("Analysis failed", f"{exc}\n\nSee terminal for details.")
                root.destroy()
            except Exception:
                pass
